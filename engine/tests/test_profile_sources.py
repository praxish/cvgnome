# SPDX-License-Identifier: MPL-2.0
from __future__ import annotations

from contextlib import closing
from copy import deepcopy
import hashlib
from io import BytesIO
import json
import os
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch
import uuid
from pathlib import Path
from typing import Any, Sequence

from docx import Document


ENGINE_SRC = Path(__file__).resolve().parents[1] / "src"
sys.path.insert(0, str(ENGINE_SRC))

import cvgnome_engine.profile_sources as profile_sources  # noqa: E402
from cvgnome_engine.profile_sources import (  # noqa: E402
    ProfileSourceWorkflowError,
    commit_profile_sources,
    discard_profile_sources,
    preview_profile_sources,
)
from cvgnome_engine.source_ingest import SourceParserError  # noqa: E402
from cvgnome_engine.resume import (  # noqa: E402
    is_baseline_resume_renderable,
    render_docx_bytes,
    render_pdf_bytes,
)
from cvgnome_engine.storage import VaultError, latest_profile, save_profile  # noqa: E402


PROFILE = {
    "basics": {
        "name": "Ada Lovelace",
        "label": "Analytics Engineer",
        "email": "ada@example.test",
        "summary": "Builds trustworthy analytics systems for consequential decisions.",
    },
    "work": [
        {
            "name": "Analytical Engines",
            "position": "Lead Engineer",
            "startDate": "2022-01",
            "highlights": ["Reduced reporting time by 40 percent."],
        }
    ],
    "education": [
        {
            "institution": "University of London",
            "studyType": "BSc",
            "area": "Mathematics",
        }
    ],
    "skills": [{"name": "Data", "keywords": ["Python", "SQL"]}],
}

FOUR_FORMAT_SOURCES = (
    (
        "private/canonical-profile.json",
        "json",
        json.dumps(PROFILE, sort_keys=True, separators=(",", ":")).encode("utf-8"),
    ),
    (
        "private/work-preferences.md",
        "md",
        (
            b"# Work Preferences\n\nRemote preference: Remote-first\n"
            b"Target roles: Staff Analytics Engineer\n"
        ),
    ),
    (
        "private/impact-evidence.txt",
        "txt",
        b"Achievement: Reduced monthly reporting time by 40 percent for 120 users.\n",
    ),
    (
        "private/skills-inventory.csv",
        "csv",
        b"Category,Value\nLanguages,Python\nDatabases,SQL\n",
    ),
)


def _conventional_docx_resume() -> bytes:
    document = Document()
    document.add_heading("Ada Lovelace", level=0)
    document.add_paragraph("ada@example.test")
    document.add_heading("Experience", level=1)
    document.add_paragraph("Lead Engineer | Analytical Engines")
    document.add_paragraph("January 2022 – Present")
    document.add_paragraph(
        "Reduced reporting time by 40 percent.",
        style="List Bullet",
    )
    document.add_heading("Skills", level=1)
    document.add_paragraph("Python, SQL")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class ProfileSourcesTests(unittest.TestCase):
    def _stage_sources(
        self,
        data_dir: Path,
        sources: Sequence[tuple[str, str, bytes]],
        *,
        scan_id: str | None = None,
        skipped: int = 0,
        scan_issues: dict[str, int] | None = None,
    ) -> dict[str, Any]:
        scan_id = scan_id or str(uuid.uuid4())
        staging = data_dir / "imports" / "staging" / scan_id
        staging.mkdir(parents=True, exist_ok=True)
        manifest: list[dict[str, Any]] = []
        for ordinal, (display_name, source_format, content) in enumerate(sources):
            path = staging / f"{ordinal:04d}.{source_format}"
            path.write_bytes(content)
            manifest.append(
                {
                    "ordinal": ordinal,
                    "managed_relative_path": path.relative_to(data_dir).as_posix(),
                    "display_name": display_name,
                    "format": source_format,
                    "byte_size": len(content),
                    "checksum_sha256": hashlib.sha256(content).hexdigest(),
                }
            )
        return {
            "scan_id": scan_id,
            "sources": manifest,
            "file_counts": {
                "discovered": len(sources) + skipped,
                "staged": len(sources),
                "skipped": skipped,
            },
            "scan_issues": scan_issues or {},
        }

    def _assert_manifest_invalid(self, data_dir: Path, params: dict[str, Any]) -> None:
        with self.assertRaises(ProfileSourceWorkflowError) as raised:
            preview_profile_sources(data_dir, params)
        self.assertEqual(raised.exception.code, "profile_source_manifest_invalid")
        self.assertEqual(str(raised.exception), "The source scan manifest is invalid.")

    def _assert_no_source_details(
        self,
        result: dict[str, Any],
        *,
        params: dict[str, Any],
        secret_values: Sequence[str] = (),
    ) -> None:
        serialized = json.dumps(result, ensure_ascii=False, sort_keys=True)
        for source in params["sources"]:
            self.assertNotIn(source["managed_relative_path"], serialized)
            self.assertNotIn(source["display_name"], serialized)
            self.assertNotIn(source["checksum_sha256"], serialized)
        for secret in secret_values:
            self.assertNotIn(secret, serialized)

    def test_native_manifest_rejects_untrusted_shapes_and_paths(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_dir = Path(directory)
            params = self._stage_sources(
                data_dir,
                [("resume.md", "md", b"# Ada Lovelace\n\n## Experience\n")],
            )

            invalid_cases: list[dict[str, Any]] = []

            noncanonical_id = deepcopy(params)
            noncanonical_id["scan_id"] = params["scan_id"].upper()
            invalid_cases.append(noncanonical_id)

            noncontiguous = deepcopy(params)
            noncontiguous["sources"][0]["ordinal"] = 1
            invalid_cases.append(noncontiguous)

            boolean_ordinal = deepcopy(params)
            boolean_ordinal["sources"][0]["ordinal"] = False
            invalid_cases.append(boolean_ordinal)

            traversal = deepcopy(params)
            traversal["sources"][0]["managed_relative_path"] = "../outside.md"
            invalid_cases.append(traversal)

            wrong_staging_name = deepcopy(params)
            wrong_staging_name["sources"][0]["managed_relative_path"] = (
                f"imports/staging/{params['scan_id']}/private-name.md"
            )
            invalid_cases.append(wrong_staging_name)

            unsupported = deepcopy(params)
            unsupported["sources"][0]["format"] = "html"
            invalid_cases.append(unsupported)

            boolean_size = deepcopy(params)
            boolean_size["sources"][0]["byte_size"] = True
            invalid_cases.append(boolean_size)

            uppercase_hash = deepcopy(params)
            uppercase_hash["sources"][0]["checksum_sha256"] = uppercase_hash["sources"][0][
                "checksum_sha256"
            ].upper()
            invalid_cases.append(uppercase_hash)

            inconsistent_counts = deepcopy(params)
            inconsistent_counts["file_counts"]["staged"] = 0
            invalid_cases.append(inconsistent_counts)

            impossible_discovered_count = deepcopy(params)
            impossible_discovered_count["file_counts"]["discovered"] = 0
            invalid_cases.append(impossible_discovered_count)

            for invalid in invalid_cases:
                with self.subTest(invalid=invalid):
                    self._assert_manifest_invalid(data_dir, invalid)

            staged = data_dir / params["sources"][0]["managed_relative_path"]
            staged.write_bytes(b"# Eve Mallory\n\n## Experience\n")
            with self.assertRaises(ProfileSourceWorkflowError) as raised:
                preview_profile_sources(data_dir, params)
            self.assertEqual(raised.exception.code, "profile_source_input_changed")
            self.assertEqual(
                str(raised.exception),
                "A staged source changed or became unavailable. Scan the folder again.",
            )

    @unittest.skipIf(os.name == "nt", "POSIX symlink semantics are unavailable")
    def test_native_manifest_rejects_symlink_substitution(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_dir = Path(directory)
            params = self._stage_sources(
                data_dir,
                [("resume.md", "md", b"# Ada Lovelace\n\n## Experience\n")],
            )
            staged = data_dir / params["sources"][0]["managed_relative_path"]
            outside = data_dir / "outside.md"
            outside.write_bytes(staged.read_bytes())
            staged.unlink()
            staged.symlink_to(outside)

            with self.assertRaises(ProfileSourceWorkflowError) as raised:
                preview_profile_sources(data_dir, params)
            self.assertEqual(raised.exception.code, "profile_source_input_changed")
            self.assertNotIn(str(outside), str(raised.exception))

    def test_four_text_formats_preview_then_commit_and_render(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_dir = Path(directory)
            params = self._stage_sources(data_dir, FOUR_FORMAT_SOURCES)

            self.assertIsNone(latest_profile(data_dir))
            preview = preview_profile_sources(data_dir, params)

            self.assertTrue(preview["can_build"])
            self.assertEqual(
                set(preview),
                {
                    "scan_id",
                    "expires_at_ms",
                    "base_profile",
                    "file_counts",
                    "source_counts",
                    "format_counts",
                    "warnings",
                    "can_build",
                "review_candidate_count",
                "source_review_count",
                },
            )
            self.assertEqual(
                preview["file_counts"],
                {
                    "discovered": 4,
                    "staged": 4,
                    "parsed": 4,
                    "duplicates": 0,
                    "skipped": 0,
                    "failed": 0,
                },
            )
            self.assertEqual(preview["format_counts"], {"csv": 1, "json": 1, "md": 1, "txt": 1})
            self.assertEqual(preview["source_counts"]["resume"], 1)
            self.assertIsNone(latest_profile(data_dir), "preview must not create a profile version")

            receipt = commit_profile_sources(data_dir, params["scan_id"])

            self.assertTrue(receipt["created"])
            self.assertEqual(receipt["version_number"], 1)
            self.assertEqual(receipt["profile_name"], "Ada Lovelace")
            self.assertTrue(receipt["renderable"])
            self.assertEqual(receipt["source_counts"]["parsed"], 4)
            self.assertFalse(
                (data_dir / "imports" / "staging" / params["scan_id"]).exists()
            )

            stored = latest_profile(data_dir)
            self.assertIsNotNone(stored)
            profile = stored["profile"]
            self.assertTrue(is_baseline_resume_renderable(profile))
            self.assertTrue(render_docx_bytes(profile).startswith(b"PK"))
            pdf = render_pdf_bytes(profile)
            self.assertTrue(pdf.startswith(b"%PDF-"))
            self.assertIn(b"%%EOF", pdf[-64:])

    def test_conventional_docx_layout_builds_a_renderable_profile(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_dir = Path(directory)
            params = self._stage_sources(
                data_dir,
                (("private/resume.docx", "docx", _conventional_docx_resume()),),
            )

            preview = preview_profile_sources(data_dir, params)

            self.assertTrue(preview["can_build"])
            self.assertEqual(preview["source_counts"]["resume"], 1)
            self.assertEqual(preview["warnings"], [])
            receipt = commit_profile_sources(data_dir, params["scan_id"])
            self.assertTrue(receipt["renderable"])
            profile = latest_profile(data_dir)["profile"]
            self.assertEqual(profile["basics"]["name"], "Ada Lovelace")
            self.assertEqual(profile["work"][0]["position"], "Lead Engineer")
            self.assertEqual(profile["work"][0]["name"], "Analytical Engines")
            self.assertEqual(
                profile["work"][0]["highlights"],
                ["Reduced reporting time by 40 percent."],
            )
            self.assertEqual(profile["skills"][0]["keywords"], ["Python", "SQL"])

    def test_structured_parse_budget_uses_remaining_time_and_degrades_safely(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_dir = Path(directory)
            sources = tuple(
                (
                    f"private/resume-{index}.docx",
                    "docx",
                    f"unique-structured-source-{index}".encode("utf-8"),
                )
                for index in range(8)
            )
            params = self._stage_sources(data_dir, sources)
            now = [1_000.0]
            observed_timeouts: list[float] = []

            def fake_monotonic() -> float:
                return now[0]

            def fake_structured_parse(*_args: Any, timeout_seconds: float, **_kwargs: Any):
                observed_timeouts.append(timeout_seconds)
                if len(observed_timeouts) <= 4:
                    now[0] += 15.0
                elif len(observed_timeouts) == 5:
                    now[0] += 14.75
                else:
                    now[0] += timeout_seconds
                    raise SourceParserError("parser_timeout")
                return {
                    "text": (
                        "Ada Lovelace\nada@example.test\nExperience\n"
                        "Lead Engineer | Analytical Engines\n"
                        "Reduced reporting time by 40 percent.\nSkills\nPython, SQL"
                    ),
                    "warnings": [],
                }

            with (
                patch.object(profile_sources, "_monotonic", side_effect=fake_monotonic),
                patch.object(
                    profile_sources,
                    "parse_structured_source",
                    side_effect=fake_structured_parse,
                ),
            ):
                preview = preview_profile_sources(data_dir, params)

            self.assertLess(profile_sources.SOURCE_PREVIEW_PARSE_BUDGET_SECONDS, 120.0)
            self.assertEqual(preview["file_counts"]["parsed"], 5)
            self.assertEqual(preview["file_counts"]["failed"], 3)
            self.assertEqual(len(observed_timeouts), 6)
            self.assertEqual(observed_timeouts[:5], [15.0] * 5)
            self.assertAlmostEqual(observed_timeouts[5], 0.25)
            self.assertEqual(
                [
                    warning
                    for warning in preview["warnings"]
                    if warning["code"] == "partial_parse" and warning["stage"] == "parse"
                ],
                [
                    {
                        "code": "partial_parse",
                        "stage": "parse",
                        "severity": "warning",
                        "count": 3,
                    }
                ],
            )
            self._assert_no_source_details(preview, params=params)

    def test_duplicate_and_failed_sources_are_aggregated_without_ui_leakage(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_dir = Path(directory)
            private_marker = "PRIVATE-CONTENT-MARKER-4821"
            profile = deepcopy(PROFILE)
            profile["basics"]["summary"] += f" {private_marker}"
            valid_json = json.dumps(profile, sort_keys=True).encode("utf-8")
            params = self._stage_sources(
                data_dir,
                (
                    ("private/top-secret-profile.json", "json", valid_json),
                    ("private/duplicate-secret.json", "json", valid_json),
                    ("private/unreadable-secret.json", "json", b"{not valid json"),
                ),
                skipped=2,
                scan_issues={"unsupported_file_type": 1, "file_too_large": 1},
            )

            preview = preview_profile_sources(data_dir, params)

            self.assertTrue(preview["can_build"])
            self.assertEqual(
                preview["file_counts"],
                {
                    "discovered": 5,
                    "staged": 3,
                    "parsed": 1,
                    "duplicates": 1,
                    "skipped": 2,
                    "failed": 1,
                },
            )
            warnings = {
                (item["code"], item["stage"], item["severity"]): item["count"]
                for item in preview["warnings"]
            }
            self.assertEqual(warnings[("duplicate_content", "parse", "info")], 1)
            self.assertEqual(warnings[("unreadable_document", "parse", "warning")], 1)
            self.assertEqual(warnings[("unsupported_file_type", "scan", "info")], 1)
            self.assertEqual(warnings[("file_too_large", "scan", "warning")], 1)
            self._assert_no_source_details(
                preview,
                params=params,
                secret_values=(private_marker, "{not valid json"),
            )

            receipt = commit_profile_sources(data_dir, params["scan_id"])

            self.assertTrue(receipt["renderable"])
            self.assertEqual(receipt["source_counts"], {"parsed": 1, "used": 1, "unused": 0})
            self._assert_no_source_details(
                receipt,
                params=params,
                secret_values=(private_marker, "{not valid json"),
            )
            self.assertEqual(
                set(receipt),
                {
                    "scan_id",
                    "profile_version_id",
                    "version_number",
                    "created",
                    "checksum_sha256",
                    "profile_name",
                    "renderable",
                    "source_counts",
                    "profile_counts",
                    "warnings",
                    "review_item_count",
                    "source_review_count",
                },
            )

    def test_identical_bytes_use_each_distinct_parser_contract(self) -> None:
        content = b'Name,Skills\nAda Lovelace,"Python, SQL"\n'
        source_orders = (
            (("private/a.txt", "txt", content), ("private/b.csv", "csv", content)),
            (("private/b.csv", "csv", content), ("private/a.txt", "txt", content)),
        )
        checksums: list[str] = []
        for sources in source_orders:
            with self.subTest(sources=[source[1] for source in sources]):
                with tempfile.TemporaryDirectory() as directory:
                    data_dir = Path(directory)
                    params = self._stage_sources(
                        data_dir,
                        sources,
                        scan_issues={"duplicate_content": 1},
                    )

                    preview = preview_profile_sources(data_dir, params)
                    receipt = commit_profile_sources(data_dir, params["scan_id"])

                    self.assertTrue(preview["can_build"])
                    self.assertEqual(preview["file_counts"]["parsed"], 2)
                    self.assertEqual(preview["file_counts"]["duplicates"], 0)
                    self.assertTrue(receipt["renderable"])
                    checksums.append(receipt["checksum_sha256"])
                    with closing(
                        sqlite3.connect(data_dir / "cvgnome.sqlite3")
                    ) as connection:
                        self.assertEqual(
                            connection.execute(
                                "SELECT COUNT(*) FROM sources WHERE content_addressed = 1"
                            ).fetchone()[0],
                            1,
                        )
                        self.assertEqual(
                            connection.execute(
                                "SELECT COUNT(*) FROM source_extractions"
                            ).fetchone()[0],
                            2,
                        )
                        self.assertEqual(
                            connection.execute(
                                "SELECT COUNT(*) FROM profile_version_sources"
                            ).fetchone()[0],
                            2,
                        )

        self.assertEqual(checksums[0], checksums[1])

    def test_aggregate_extraction_budget_degrades_to_partial_results(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_dir = Path(directory)
            sources = []
            for index in range(11):
                prefix = f"Private notes batch {index}\n".encode("utf-8")
                content = prefix + bytes([65 + index]) * (200_000 - len(prefix))
                sources.append((f"private/batch-{index}.txt", "txt", content))
            sources.append(
                (
                    "private/z-canonical-profile.json",
                    FOUR_FORMAT_SOURCES[0][1],
                    FOUR_FORMAT_SOURCES[0][2],
                )
            )
            params = self._stage_sources(data_dir, sources)

            preview = preview_profile_sources(data_dir, params)

            self.assertTrue(preview["can_build"])
            self.assertEqual(preview["file_counts"]["parsed"], 11)
            self.assertEqual(preview["file_counts"]["failed"], 1)
            partial = [
                warning
                for warning in preview["warnings"]
                if warning["code"] == "partial_parse" and warning["stage"] == "parse"
            ]
            self.assertEqual(
                partial,
                [
                    {
                        "code": "partial_parse",
                        "stage": "parse",
                        "severity": "warning",
                        "count": 2,
                    }
                ],
            )

            with closing(
                sqlite3.connect(data_dir / "cvgnome.sqlite3")
            ) as connection:
                extracted_chars = connection.execute(
                    "SELECT SUM(length(extracted_text)) FROM profile_source_scan_items"
                ).fetchone()[0]
            self.assertEqual(extracted_chars, 2_000_000)

            receipt = commit_profile_sources(data_dir, params["scan_id"])
            self.assertTrue(receipt["renderable"])
            self.assertEqual(receipt["source_counts"]["parsed"], 11)

    def test_exhausted_extraction_budget_is_source_order_invariant(self) -> None:
        sources: list[tuple[str, str, bytes]] = []
        for marker in "ABCDEFGHIJK":
            prefix = f"Name: Person {marker}\nSkills: Skill{marker}\n".encode("utf-8")
            content = prefix + marker.encode("ascii") * (200_000 - len(prefix))
            sources.append((f"private/{marker.lower()}.txt", "txt", content))

        results: list[tuple[dict[str, Any], dict[str, Any]]] = []
        for source_order in (sources, list(reversed(sources))):
            with tempfile.TemporaryDirectory() as directory:
                data_dir = Path(directory)
                params = self._stage_sources(data_dir, source_order)
                preview = preview_profile_sources(data_dir, params)
                receipt = commit_profile_sources(data_dir, params["scan_id"])
                results.append((preview, receipt))

        self.assertEqual(results[0][0]["file_counts"]["failed"], 1)
        self.assertEqual(results[1][0]["file_counts"]["failed"], 1)
        self.assertEqual(results[0][0]["warnings"], results[1][0]["warnings"])
        self.assertEqual(
            results[0][1]["checksum_sha256"],
            results[1][1]["checksum_sha256"],
        )

    def test_stale_base_is_refused_and_preview_can_be_discarded(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_dir = Path(directory)
            base = save_profile(
                data_dir,
                {"basics": {"name": "Ada Lovelace", "summary": "Original summary."}},
            )
            params = self._stage_sources(data_dir, FOUR_FORMAT_SOURCES)
            preview = preview_profile_sources(data_dir, params)
            self.assertEqual(preview["base_profile"]["profile_version_id"], base.id)

            replacement = save_profile(
                data_dir,
                {"basics": {"name": "Grace Hopper", "summary": "New current profile."}},
            )
            with self.assertRaises(VaultError) as raised:
                commit_profile_sources(data_dir, params["scan_id"])
            self.assertEqual(raised.exception.code, "profile_source_base_changed")
            self.assertEqual(latest_profile(data_dir)["id"], replacement.id)

            discarded = discard_profile_sources(data_dir, params["scan_id"])
            self.assertEqual(discarded, {"discarded": True})
            self.assertFalse(
                (data_dir / "imports" / "staging" / params["scan_id"]).exists()
            )

    def test_committed_scan_retry_survives_a_newer_profile_version(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_dir = Path(directory)
            params = self._stage_sources(data_dir, FOUR_FORMAT_SOURCES)
            preview_profile_sources(data_dir, params)
            first = commit_profile_sources(data_dir, params["scan_id"])
            newer_profile = deepcopy(PROFILE)
            newer_profile["basics"]["name"] = "Grace Hopper"
            newer = save_profile(data_dir, newer_profile)

            retry = commit_profile_sources(data_dir, params["scan_id"])

            self.assertEqual(retry, first)
            self.assertEqual(latest_profile(data_dir)["id"], newer.id)

    def test_discard_is_explicit_idempotent_and_never_creates_a_profile(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_dir = Path(directory)
            params = self._stage_sources(data_dir, FOUR_FORMAT_SOURCES)
            preview_profile_sources(data_dir, params)

            self.assertEqual(
                discard_profile_sources(data_dir, params["scan_id"]),
                {"discarded": True},
            )
            self.assertEqual(
                discard_profile_sources(data_dir, params["scan_id"]),
                {"discarded": False},
            )
            self.assertIsNone(latest_profile(data_dir))
            self.assertFalse(
                (data_dir / "imports" / "staging" / params["scan_id"]).exists()
            )

    def test_source_order_does_not_change_the_profile(self) -> None:
        results: list[tuple[dict[str, Any], dict[str, Any]]] = []
        orders = (FOUR_FORMAT_SOURCES, tuple(reversed(FOUR_FORMAT_SOURCES)))
        for sources in orders:
            with tempfile.TemporaryDirectory() as directory:
                data_dir = Path(directory)
                params = self._stage_sources(data_dir, sources)
                preview = preview_profile_sources(data_dir, params)
                receipt = commit_profile_sources(data_dir, params["scan_id"])
                results.append(
                    (
                        preview,
                        {
                            "receipt": receipt,
                            "profile": latest_profile(data_dir)["profile"],
                        },
                    )
                )

        first_preview, first_result = results[0]
        second_preview, second_result = results[1]
        for field in ("file_counts", "source_counts", "format_counts", "warnings", "can_build"):
            self.assertEqual(first_preview[field], second_preview[field])
        self.assertEqual(
            first_result["receipt"]["checksum_sha256"],
            second_result["receipt"]["checksum_sha256"],
        )
        self.assertEqual(first_result["profile"], second_result["profile"])


if __name__ == "__main__":
    unittest.main()
