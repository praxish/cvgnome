# SPDX-License-Identifier: MPL-2.0
"""In-memory DOCX and PDF rendering for baseline resumes."""

from __future__ import annotations

from datetime import datetime
import html
from io import BytesIO
from pathlib import Path
import re
from typing import Any
import unicodedata
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from .derive import derive_baseline_resume, projected_resume_is_renderable


_ACCENT = "185B7A"
_INK = "17212B"
_MUTED = "52606D"
_RULE = "BCC8D0"
_FIXED_DOC_TIME = datetime(2000, 1, 1, 0, 0, 0)
_ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)


def _plain(value: Any) -> str:
    cleaned = unicodedata.normalize("NFC", str(value or ""))
    cleaned = re.sub(r"[\u2010-\u2015\u2212]", "-", cleaned)
    cleaned = cleaned.replace("\u00a0", " ").replace("\u2022", "-")
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned


def _join(parts: list[str]) -> str:
    return " | ".join(part for part in parts if part)


def _format_location(value: Any) -> str:
    if not isinstance(value, dict):
        return ""
    city = _plain(value.get("city"))
    region = _plain(value.get("region"))
    country = _plain(value.get("countryCode"))
    if city and region:
        region_marker = re.sub(r"[^a-z0-9]+", "", region.casefold())
        city_tail = re.split(r"[\s,]+", city.strip())[-1]
        city_tail_marker = re.sub(r"[^a-z0-9]+", "", city_tail.casefold())
        if region_marker and city_tail_marker == region_marker:
            return city
        return f"{city}, {region}"
    return city or region or country


def _format_date(value: Any) -> str:
    cleaned = _plain(value)
    if not cleaned:
        return ""
    if cleaned.casefold() in {"current", "now", "present", "ongoing"}:
        return "Present"
    match = re.fullmatch(r"(\d{4})-(\d{1,2})(?:-\d{1,2})?", cleaned)
    if not match:
        return cleaned
    month = int(match.group(2))
    if not 1 <= month <= 12:
        return cleaned
    months = (
        "",
        "Jan",
        "Feb",
        "Mar",
        "Apr",
        "May",
        "Jun",
        "Jul",
        "Aug",
        "Sep",
        "Oct",
        "Nov",
        "Dec",
    )
    return f"{months[month]} {match.group(1)}"


def _format_dates(start: Any, end: Any) -> str:
    start_text = _format_date(start)
    end_text = _format_date(end)
    if start_text and end_text:
        return f"{start_text} - {end_text}"
    if start_text:
        return f"{start_text} - Present"
    return end_text


def _profile_links(basics: dict[str, Any]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for raw in [basics.get("url"), *(basics.get("profiles") or [])]:
        if isinstance(raw, dict):
            url = _plain(raw.get("url"))
            label = _plain(raw.get("username")) or _plain(raw.get("network"))
            value = _join([label, url]) if label and not url.endswith(label) else (url or label)
        else:
            value = _plain(raw)
        marker = value.casefold()
        if value and marker not in seen:
            seen.add(marker)
            result.append(re.sub(r"^[a-z]+://(?:www\.)?", "", value, flags=re.I).rstrip("/"))
    return result


def _canonicalize_docx(raw: bytes) -> bytes:
    """Normalize ZIP metadata so identical resumes produce identical bytes."""

    source = BytesIO(raw)
    target = BytesIO()
    with ZipFile(source, "r") as archive, ZipFile(
        target,
        "w",
        compression=ZIP_DEFLATED,
        compresslevel=9,
    ) as output:
        for name in sorted(archive.namelist()):
            info = ZipInfo(name, date_time=_ZIP_EPOCH)
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            info.flag_bits = 0
            output.writestr(info, archive.read(name), compress_type=ZIP_DEFLATED, compresslevel=9)
    return target.getvalue()


def _validate_renderable(resume: dict[str, Any]) -> None:
    basics = resume.get("basics") if isinstance(resume.get("basics"), dict) else {}
    if not _plain(basics.get("name")):
        raise ValueError("a resume requires basics.name")
    if not projected_resume_is_renderable(resume):
        raise ValueError("a resume requires a summary or a nonempty structured section")


def _docx_set_font(style: Any, font_name: str, size: Any, qn: Any) -> None:
    style.font.name = font_name
    style.font.size = size
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), font_name)


def render_docx_bytes(resume: dict[str, Any]) -> bytes:
    """Render a resume to a deterministic, standards-compliant DOCX payload."""

    if not isinstance(resume, dict):
        raise TypeError("resume must be a dictionary")
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_TAB_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:  # pragma: no cover - exercised without optional deps
        raise RuntimeError(
            "python-docx is required for DOCX rendering; install `python-docx>=1.1`."
        ) from exc

    projected = derive_baseline_resume(resume)
    _validate_renderable(projected)
    basics = projected.get("basics") if isinstance(projected.get("basics"), dict) else {}
    doc = Document()
    doc.core_properties.title = _join([_plain(basics.get("name")), "Resume"]) or "Resume"
    doc.core_properties.subject = "Resume"
    doc.core_properties.author = "CVGnome"
    doc.core_properties.created = _FIXED_DOC_TIME
    doc.core_properties.modified = _FIXED_DOC_TIME
    doc.core_properties.last_modified_by = "CVGnome"
    doc.core_properties.revision = 1

    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    content_width = 8.5 - 0.72 - 0.72

    normal = doc.styles["Normal"]
    _docx_set_font(normal, "Arial", Pt(9.5), qn)
    normal.font.color.rgb = RGBColor.from_string(_INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.05

    heading = doc.styles["Heading 1"]
    _docx_set_font(heading, "Arial", Pt(9.5), qn)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(_ACCENT)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.keep_together = True

    bullet = doc.styles["List Bullet"]
    _docx_set_font(bullet, "Arial", Pt(9.25), qn)
    bullet.font.color.rgb = RGBColor.from_string(_INK)
    bullet.paragraph_format.left_indent = Inches(0.22)
    bullet.paragraph_format.first_line_indent = Inches(-0.14)
    bullet.paragraph_format.space_before = Pt(0)
    bullet.paragraph_format.space_after = Pt(1.5)
    bullet.paragraph_format.line_spacing = 1.02

    if "CVGnome Role" not in [style.name for style in doc.styles]:
        role_style = doc.styles.add_style("CVGnome Role", WD_STYLE_TYPE.PARAGRAPH)
    else:  # pragma: no cover - default template does not define it
        role_style = doc.styles["CVGnome Role"]
    _docx_set_font(role_style, "Arial", Pt(9.75), qn)
    role_style.font.bold = True
    role_style.font.color.rgb = RGBColor.from_string(_INK)
    role_style.paragraph_format.space_before = Pt(4)
    role_style.paragraph_format.space_after = Pt(1)
    role_style.paragraph_format.keep_with_next = True

    def add_section_title(title: str) -> None:
        paragraph = doc.add_paragraph(style="Heading 1")
        paragraph.add_run(title.upper())
        properties = paragraph._p.get_or_add_pPr()
        borders = properties.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            properties.append(borders)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "5")
        bottom.set(qn("w:space"), "2")
        bottom.set(qn("w:color"), _RULE)
        borders.append(bottom)

    def add_line(
        text: str,
        *,
        bold: bool = False,
        italic: bool = False,
        color: str | None = None,
        size: float | None = None,
        keep_with_next: bool = False,
    ) -> Any:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.keep_with_next = keep_with_next
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = "Arial"
        if size is not None:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        return paragraph

    def add_role_header(left: str, right: str) -> None:
        paragraph = doc.add_paragraph(style="CVGnome Role")
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Inches(content_width), WD_TAB_ALIGNMENT.RIGHT
        )
        paragraph.add_run(left)
        if right:
            right_run = paragraph.add_run(f"\t{right}")
            right_run.bold = False
            right_run.font.size = Pt(8.5)
            right_run.font.color.rgb = RGBColor.from_string(_MUTED)

    name = _plain(basics.get("name"))
    label = _plain(basics.get("label"))
    if name:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(name)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(22)
        run.font.color.rgb = RGBColor.from_string(_INK)
    if label:
        add_line(label, bold=True, color=_ACCENT, size=10.5)
    contact = _join(
        [
            _format_location(basics.get("location")),
            _plain(basics.get("phone")),
            _plain(basics.get("email")),
            *_profile_links(basics),
        ]
    )
    if contact:
        add_line(contact, color=_MUTED, size=8.5)

    summary = _plain(basics.get("summary"))
    if summary:
        add_section_title("Profile")
        add_line(summary)

    skills = projected.get("skills") or []
    if skills:
        add_section_title("Skills")
        for item in skills:
            if not isinstance(item, dict):
                continue
            skill_name = _plain(item.get("name"))
            keywords = ", ".join(_plain(value) for value in item.get("keywords") or [] if _plain(value))
            level = _plain(item.get("level"))
            line = _join([keywords, level])
            paragraph = doc.add_paragraph()
            if skill_name:
                paragraph.add_run(f"{skill_name}: ").bold = True
            paragraph.add_run(line or skill_name)

    work = projected.get("work") or []
    if work:
        add_section_title("Experience")
        for item in work:
            if not isinstance(item, dict):
                continue
            left = _join([_plain(item.get("name")), _plain(item.get("position"))])
            right = _join(
                [
                    _format_dates(item.get("startDate"), item.get("endDate")),
                    _plain(item.get("location")),
                ]
            )
            if left:
                add_role_header(left, right)
            role_summary = _plain(item.get("summary"))
            if role_summary:
                add_line(role_summary, italic=True, color=_MUTED, size=9)
            for value in item.get("highlights") or []:
                text = _plain(value)
                if text:
                    paragraph = doc.add_paragraph(text, style="List Bullet")
                    paragraph.paragraph_format.keep_together = True

    projects = projected.get("projects") or []
    if projects:
        add_section_title("Projects")
        for item in projects:
            if not isinstance(item, dict):
                continue
            add_role_header(
                _plain(item.get("name")),
                _format_dates(item.get("startDate"), item.get("endDate")),
            )
            description = _plain(item.get("description"))
            if description:
                add_line(description)
            for value in item.get("highlights") or []:
                if text := _plain(value):
                    doc.add_paragraph(text, style="List Bullet")

    education = projected.get("education") or []
    if education:
        add_section_title("Education")
        for item in education:
            if not isinstance(item, dict):
                continue
            credential = _join([_plain(item.get("studyType")), _plain(item.get("area"))])
            add_role_header(
                _join([_plain(item.get("institution")), credential]),
                _format_dates(item.get("startDate"), item.get("endDate")),
            )
            details = _join([_plain(item.get("score")), _plain(item.get("summary"))])
            if details:
                add_line(details, color=_MUTED, size=9)

    simple_sections = (
        ("certificates", "Certifications", "name", ("issuer", "date")),
        ("awards", "Awards", "title", ("awarder", "date")),
        ("publications", "Publications", "name", ("publisher", "releaseDate")),
        ("volunteer", "Volunteer", "organization", ("position", "startDate", "endDate")),
        ("languages", "Languages", "language", ("fluency",)),
        ("interests", "Interests", "name", ()),
        ("references", "References", "name", ("reference",)),
    )
    for key, title, primary_key, detail_keys in simple_sections:
        items = projected.get(key) or []
        rows: list[str] = []
        for item in items:
            if not isinstance(item, dict):
                continue
            row = _join([_plain(item.get(primary_key)), *[_plain(item.get(field)) for field in detail_keys]])
            if row:
                rows.append(row)
        if rows:
            add_section_title(title)
            for row in rows:
                add_line(row)

    if len(doc.paragraphs) == 0:
        add_line("Resume", bold=True, size=20)

    buffer = BytesIO()
    doc.save(buffer)
    return _canonicalize_docx(buffer.getvalue())


def _pdf_font_names() -> tuple[str, str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import reportlab

    regular = "CVGnomeSans"
    bold = "CVGnomeSans-Bold"
    italic = "CVGnomeSans-Italic"
    registered = set(pdfmetrics.getRegisteredFontNames())
    font_dir = Path(reportlab.__file__).resolve().parent / "fonts"
    paths = {
        regular: font_dir / "Vera.ttf",
        bold: font_dir / "VeraBd.ttf",
        italic: font_dir / "VeraIt.ttf",
    }
    if all(path.exists() for path in paths.values()):
        for name, path in paths.items():
            if name not in registered:
                pdfmetrics.registerFont(TTFont(name, str(path)))
        pdfmetrics.registerFontFamily(
            regular,
            normal=regular,
            bold=bold,
            italic=italic,
            boldItalic=bold,
        )
        return regular, bold, italic
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"


def _assert_pdf_font_coverage(resume: dict[str, Any], font_name: str) -> None:
    """Raise a useful error instead of silently emitting missing-glyph boxes."""

    from reportlab.pdfbase import pdfmetrics

    face = getattr(pdfmetrics.getFont(font_name), "face", None)
    char_to_glyph = getattr(face, "charToGlyph", None)

    def walk(value: Any):
        if isinstance(value, dict):
            for child in value.values():
                yield from walk(child)
        elif isinstance(value, list):
            for child in value:
                yield from walk(child)
        elif isinstance(value, str):
            yield _plain(value)

    for text in walk(resume):
        for character in text:
            if character.isspace() or ord(character) < 32:
                continue
            if isinstance(char_to_glyph, dict):
                supported = ord(character) in char_to_glyph
            else:
                try:
                    character.encode("cp1252")
                except UnicodeEncodeError:
                    supported = False
                else:
                    supported = True
            if not supported:
                raise ValueError(
                    "PDF font coverage does not support character "
                    f"{character!r} (U+{ord(character):04X})"
                )


def render_pdf_bytes(resume: dict[str, Any]) -> bytes:
    """Render a resume to deterministic, searchable PDF bytes."""

    if not isinstance(resume, dict):
        raise TypeError("resume must be a dictionary")
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfgen import canvas as canvas_module
        from reportlab.platypus import (
            CondPageBreak,
            HRFlowable,
            KeepTogether,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
        )
    except ImportError as exc:  # pragma: no cover - exercised without optional deps
        raise RuntimeError(
            "ReportLab is required for PDF rendering; install `reportlab>=4.0`."
        ) from exc

    projected = derive_baseline_resume(resume)
    _validate_renderable(projected)
    basics = projected.get("basics") if isinstance(projected.get("basics"), dict) else {}
    buffer = BytesIO()
    regular_font, bold_font, italic_font = _pdf_font_names()
    _assert_pdf_font_coverage(projected, regular_font)
    ink = colors.HexColor(f"#{_INK}")
    muted = colors.HexColor(f"#{_MUTED}")
    accent = colors.HexColor(f"#{_ACCENT}")
    rule = colors.HexColor(f"#{_RULE}")

    doc = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        leftMargin=44,
        rightMargin=44,
        topMargin=38,
        bottomMargin=36,
        pageCompression=1,
        allowSplitting=1,
        title=_join([_plain(basics.get("name")), "Resume"]) or "Resume",
        author=_plain(basics.get("name")) or "CVGnome",
        subject="Resume",
        creator="CVGnome",
    )

    name_style = ParagraphStyle(
        "cv-name",
        fontName=bold_font,
        fontSize=23,
        leading=25,
        textColor=ink,
        spaceAfter=0,
    )
    label_style = ParagraphStyle(
        "cv-label",
        fontName=bold_font,
        fontSize=10.5,
        leading=12.5,
        textColor=accent,
        spaceAfter=1,
    )
    contact_style = ParagraphStyle(
        "cv-contact",
        fontName=regular_font,
        fontSize=8.2,
        leading=10,
        textColor=muted,
        spaceAfter=0,
    )
    section_style = ParagraphStyle(
        "cv-section",
        fontName=bold_font,
        fontSize=8.8,
        leading=10,
        textColor=accent,
        spaceBefore=0,
        spaceAfter=0,
        keepWithNext=True,
    )
    body_style = ParagraphStyle(
        "cv-body",
        fontName=regular_font,
        fontSize=8.9,
        leading=11.2,
        textColor=ink,
        alignment=TA_LEFT,
        spaceAfter=0,
        splitLongWords=True,
    )
    role_style = ParagraphStyle(
        "cv-role",
        parent=body_style,
        fontName=bold_font,
        fontSize=9.7,
        leading=11.5,
    )
    meta_style = ParagraphStyle(
        "cv-meta",
        parent=body_style,
        fontSize=8.1,
        leading=9.5,
        textColor=muted,
    )
    summary_style = ParagraphStyle(
        "cv-summary",
        parent=body_style,
        fontName=italic_font,
        textColor=muted,
    )
    bullet_style = ParagraphStyle(
        "cv-bullet",
        parent=body_style,
        leftIndent=11,
        firstLineIndent=0,
        bulletIndent=1,
        bulletFontName=regular_font,
        bulletFontSize=8.9,
        spaceBefore=0.5,
    )

    story: list[Any] = []

    def markup(value: Any) -> str:
        return html.escape(_plain(value), quote=True)

    def add_section(title: str, content: list[Any]) -> None:
        if not content:
            return
        # A small conditional guard prevents orphaned section headers without
        # nesting KeepTogether blocks. Nested blocks cause ReportLab to promote
        # otherwise tiny sections onto their own pages.
        story.append(CondPageBreak(72))
        header = [
            Spacer(1, 7),
            Paragraph(markup(title.upper()), section_style),
            HRFlowable(width="100%", thickness=0.55, color=rule, spaceBefore=2, spaceAfter=3.5),
        ]
        story.extend(header)
        story.extend(content)

    name = _plain(basics.get("name"))
    label = _plain(basics.get("label"))
    if name:
        story.append(Paragraph(markup(name), name_style))
    if label:
        story.append(Paragraph(markup(label), label_style))
    contact = _join(
        [
            _format_location(basics.get("location")),
            _plain(basics.get("phone")),
            _plain(basics.get("email")),
            *_profile_links(basics),
        ]
    )
    if contact:
        story.append(Paragraph(markup(contact), contact_style))

    summary = _plain(basics.get("summary"))
    if summary:
        add_section("Profile", [Paragraph(markup(summary), body_style)])

    skill_content: list[Any] = []
    for item in projected.get("skills") or []:
        if not isinstance(item, dict):
            continue
        group = _plain(item.get("name"))
        keywords = ", ".join(_plain(value) for value in item.get("keywords") or [] if _plain(value))
        level = _plain(item.get("level"))
        detail = _join([keywords, level])
        if group and detail:
            row = f"<b>{markup(group)}:</b> {markup(detail)}"
        else:
            row = markup(group or detail)
        if row:
            skill_content.append(Paragraph(row, body_style))
    add_section("Skills", skill_content)

    experience_content: list[Any] = []
    for item in projected.get("work") or []:
        if not isinstance(item, dict):
            continue
        left = _join([_plain(item.get("name")), _plain(item.get("position"))])
        if not left:
            continue
        if experience_content:
            experience_content.append(Spacer(1, 4))
        role_parts: list[Any] = [Paragraph(markup(left), role_style)]
        metadata = _join(
            [
                _format_dates(item.get("startDate"), item.get("endDate")),
                _plain(item.get("location")),
            ]
        )
        if metadata:
            role_parts.append(Paragraph(markup(metadata), meta_style))
        role_summary = _plain(item.get("summary"))
        if role_summary:
            role_parts.append(Paragraph(markup(role_summary), summary_style))
        highlights = [_plain(value) for value in item.get("highlights") or [] if _plain(value)]
        if highlights:
            role_parts.append(Paragraph(markup(highlights[0]), bullet_style, bulletText="-"))
        experience_content.append(KeepTogether(role_parts))
        for value in highlights[1:]:
            experience_content.append(Paragraph(markup(value), bullet_style, bulletText="-"))
    add_section("Experience", experience_content)

    project_content: list[Any] = []
    for item in projected.get("projects") or []:
        if not isinstance(item, dict):
            continue
        if project_content:
            project_content.append(Spacer(1, 3))
        title = _plain(item.get("name"))
        dates = _format_dates(item.get("startDate"), item.get("endDate"))
        parts = [Paragraph(markup(title), role_style)] if title else []
        if dates:
            parts.append(Paragraph(markup(dates), meta_style))
        description = _plain(item.get("description"))
        if description:
            parts.append(Paragraph(markup(description), body_style))
        if parts:
            project_content.append(KeepTogether(parts))
    add_section("Projects", project_content)

    education_content: list[Any] = []
    for item in projected.get("education") or []:
        if not isinstance(item, dict):
            continue
        credential = _join([_plain(item.get("studyType")), _plain(item.get("area"))])
        title = _join([_plain(item.get("institution")), credential])
        if not title:
            continue
        parts = [Paragraph(markup(title), role_style)]
        dates = _format_dates(item.get("startDate"), item.get("endDate"))
        details = _join([dates, _plain(item.get("score")), _plain(item.get("summary"))])
        if details:
            parts.append(Paragraph(markup(details), meta_style))
        education_content.append(KeepTogether(parts))
    add_section("Education", education_content)

    simple_sections = (
        ("certificates", "Certifications", "name", ("issuer", "date")),
        ("awards", "Awards", "title", ("awarder", "date")),
        ("publications", "Publications", "name", ("publisher", "releaseDate")),
        ("volunteer", "Volunteer", "organization", ("position", "startDate", "endDate")),
        ("languages", "Languages", "language", ("fluency",)),
        ("interests", "Interests", "name", ()),
        ("references", "References", "name", ("reference",)),
    )
    for key, title, primary, details in simple_sections:
        content: list[Any] = []
        for item in projected.get(key) or []:
            if not isinstance(item, dict):
                continue
            row = _join([_plain(item.get(primary)), *[_plain(item.get(field)) for field in details]])
            if row:
                content.append(Paragraph(markup(row), body_style))
        add_section(title, content)

    if not story:
        story.append(Paragraph("Resume", name_style))

    candidate_name = name or "CVGnome"

    def on_page(canvas: Any, document: Any) -> None:
        canvas.setTitle(_join([candidate_name, "Resume"]))
        canvas.setAuthor(candidate_name)
        canvas.setSubject("Resume")
        canvas.setCreator("CVGnome")
        canvas.saveState()
        canvas.setFont(regular_font, 7)
        canvas.setFillColor(muted)
        canvas.drawRightString(LETTER[0] - 44, 20, f"Page {document.page}")
        canvas.restoreState()

    class InvariantCanvas(canvas_module.Canvas):
        def __init__(self, *args: Any, **kwargs: Any) -> None:
            kwargs["invariant"] = 1
            super().__init__(*args, **kwargs)

    doc.build(
        story,
        onFirstPage=on_page,
        onLaterPages=on_page,
        canvasmaker=InvariantCanvas,
    )
    payload = buffer.getvalue()
    if not payload.startswith(b"%PDF-") or b"%%EOF" not in payload[-64:]:
        raise RuntimeError("PDF renderer produced an invalid payload")
    return payload
